"""Build the NDMA Digital Twin project tracker (layout mirrors the KONU tracker)."""
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
PROJECT_START = date(2026, 10, 12)

# Holidays used both by the Python scheduler and by NETWORKDAYS in the sheet.
HOLIDAYS = [
    (date(2026, 10, 20), "Dussehra"),
    (date(2026, 12, 25), "Christmas"),
    (date(2027, 1, 26), "Republic Day"),
    (date(2027, 3, 26), "Good Friday"),
]
HOL = {d for d, _ in HOLIDAYS}

D_MEET = "Needs: NDMA meeting slot for walkthrough."
D_HOST = "Needs: Hosting decision (H0)."

# Each row: id, module, category, task, subtask, description, done-when, duration (working days),
# deps (ids), status, needs (external blocker text), comments, artifacts, extra options.
R = []


def t(id_, mod, cat, task, sub, desc, dod, dur, deps=(), status="Not Started", needs="",
      comment="", artifact="", start=None, not_before=None):
    R.append(dict(id=id_, mod=mod, cat=cat, task=task, sub=sub, desc=desc, dod=dod, dur=dur,
                  deps=list(deps), status=status, needs=needs, comment=comment,
                  artifact=artifact, start=start, not_before=not_before))


# ---------------------------------------------------------------- A
A = "A – Programme Setup & Governance"
c = "1. Project Initiation"
t("P0", A, c, "Solution Proposal", "Draft AI-integrated Digital Twin proposal",
  "Write the solution proposal: problem statement, five-layer solution, dashboard features, usage example, data sources, risks and expected impact.",
  "Done when: proposal (PDF + DOCX) is complete and committed to the repo.",
  10, status="Done", start=date(2026, 9, 21),
  artifact="NDMA_Digital_Twin_Proposal.pdf; NDMA_Digital_Twin_Proposal (1).docx")
t("P1", A, c, "Solution Proposal", "Proposal walkthrough and sign-off with NDMA",
  "Present the proposal to NDMA, capture feedback and get written go-ahead for Phase 1.",
  "Done when: NDMA sign-off (email/minutes) received and feedback logged.",
  4, ["P0"], needs=D_MEET)
t("P2", A, c, "Scope & Pilot", "Select Phase 1 pilot district",
  "Pick one well-mapped, flood-prone district with CWC gauges, good DEM coverage and a willing DEOC (mitigates the terrain-resolution risk).",
  "Done when: pilot district agreed in writing with a scored selection matrix.",
  5, ["P1"], needs="Needs: NDMA / State SDMA nomination of candidate districts.")
t("P3", A, c, "Scope & Pilot", "Lock Phase 1 hazard and scope",
  "Phase 1 = riverine flood in the pilot district. Landslide, cyclone and wider sensor/crowd inputs move to Phase 2; multi-district scale-up is Phase 3.",
  "Done when: in/out-of-scope list signed off by NDMA.",
  3, ["P2"])
t("P4", A, c, "Scope & Pilot", "User roles and operational workflows",
  "Interview NDMA, SEOC, DEOC and NDRF/SDRF users; map how they plan shelters, routes and deployments today.",
  "Done when: personas, current-vs-future workflow maps and a prioritised needs list are documented.",
  6, ["P1"])
t("P5", A, c, "Scope & Pilot", "Success metrics and acceptance bar",
  "Agree measurable targets: flood-extent accuracy (IoU/CSI), forecast lead time, route recompute time, dashboard load time.",
  "Done when: metric targets are agreed with NDMA and referenced by validation tasks.",
  4, ["P3", "P4"])
c = "2. Data Agreements & Governance"
t("D1", A, c, "Data-Sharing Agreements", "IMD data-sharing agreement",
  "Formal agreement for rainfall nowcasts, district forecasts and cyclone bulletins (API/FTP access).",
  "Done when: signed agreement and working credentials/endpoint for the pilot district.",
  20, ["P2"], needs="Needs: IMD nodal officer; NDMA letter of request.")
t("D2", A, c, "Data-Sharing Agreements", "CWC data-sharing agreement",
  "Agreement for real-time river gauge levels and flood forecasts.",
  "Done when: signed agreement and live gauge feed access for pilot basins.",
  20, ["P2"], needs="Needs: CWC nodal officer; NDMA letter of request.")
t("D3", A, c, "Data-Sharing Agreements", "ISRO / NRSC satellite data access",
  "Access to optical and SAR imagery and flood maps (Bhuvan / NDEM) for the pilot district.",
  "Done when: imagery access confirmed and a sample scene downloaded.",
  25, ["P2"], needs="Needs: NRSC approval via NDMA.")
t("D4", A, c, "Data-Sharing Agreements", "Survey of India high-resolution DEM licence",
  "Licence for high-resolution elevation data — the base terrain of the 3D twin.",
  "Done when: licence signed and DEM tiles for the pilot district received.",
  20, ["P2"], needs="Needs: SoI licence approval.")
t("D5", A, c, "Data-Sharing Agreements", "Census / LGD data access",
  "Population, household and administrative boundary data for impact and shelter sizing.",
  "Done when: ward/village-level population and LGD boundaries for the pilot district are available.",
  10, ["P2"])
t("D6", A, c, "Data-Sharing Agreements", "Google Maps Platform licence and billing",
  "Enable Roads, Places and Traffic APIs with a government billing account and usage caps.",
  "Done when: API keys issued with quotas and a cost estimate approved.",
  8, ["P1"], needs="Needs: budget approval for API usage.")
t("D7", A, c, "Data-Sharing Agreements", "Commercial high-resolution imagery procurement",
  "Procure fine-resolution imagery for the pilot district to sharpen terrain and infrastructure detail.",
  "Done when: purchase order issued and imagery delivered.",
  30, ["P2"], needs="Needs: procurement approval (GeM).", comment="Pilot district only in Phase 1.")
t("G1", A, c, "Privacy & Data Governance", "Data governance and privacy policy",
  "Policy covering citizen location/photo data and social media signals under the DPDP Act 2023: purpose limitation, minimisation, retention, access.",
  "Done when: policy approved by NDMA legal and published internally.",
  15, ["P1"], needs="Needs: NDMA legal review.",
  comment="Runs alongside the technical build (proposal risk: data privacy & consent).")
t("G2", A, c, "Privacy & Data Governance", "Data classification and retention rules",
  "Classify each data source (public / restricted / personal) and set retention periods.",
  "Done when: every ingested source has a class and retention rule applied in the pipeline config.",
  5, ["G1"])
t("G3", A, c, "Privacy & Data Governance", "Consent flow for citizen ground reports",
  "Consent text and opt-in flow for WhatsApp / SACHET reports (photo + location).",
  "Done when: consent text approved and recorded with every citizen report.",
  5, ["G1"])

# ---------------------------------------------------------------- B
B = "B – Unified Data Ingestion (Layer 1)"
c = "3. Ingestion Foundation"
t("I1", B, c, "Common Data Model", "Canonical schema for all feeds",
  "One schema for every source: timestamp (UTC + IST), geometry, source, unit, confidence and raw-payload reference.",
  "Done when: schema documented and validated against samples from every Phase 1 source.",
  8, ["P3"])
t("I2", B, c, "Ingestion Framework", "Connector framework and scheduler",
  "Reusable connector framework with scheduling, retries, back-fill and idempotent loads so feeds stay fresh without manual effort.",
  "Done when: scheduled runs complete, resume after interruption and never run twice at once.",
  10, ["I1", "H1"])
t("I3", B, c, "Ingestion Framework", "Spatial data store (PostGIS + object storage)",
  "Store vector layers, time-series and rasters/imagery in one spatially indexed store.",
  "Done when: store is provisioned in dev/test with spatial indexes and backup.",
  8, ["I1", "H1"])
t("I4", B, c, "Ingestion Framework", "Pipeline status view (last run, errors, owner)",
  "Simple view showing, per source, when it last succeeded, what failed and who owns it.",
  "Done when: operations can spot a stale/failed feed without engineering help.",
  5, ["I2"])
c = "4. Agency Feeds"
t("I5", B, c, "IMD", "IMD connector (rainfall, forecasts, cyclone bulletins)",
  "Pull nowcasts, district rainfall forecasts and bulletins into the canonical schema.",
  "Done when: IMD data for the pilot district lands every cycle with < 1 cycle lag.",
  8, ["I2", "D1"])
t("I6", B, c, "CWC", "CWC connector (gauge levels, flood forecasts)",
  "Pull river gauge levels, danger/warning marks and flood forecasts for the pilot basins.",
  "Done when: all pilot-basin gauges update on schedule with danger marks attached.",
  8, ["I2", "D2"])
t("I7", B, c, "ISRO / NRSC", "Satellite imagery connector (optical + SAR)",
  "Ingest optical and SAR scenes and NRSC flood-extent products for change detection.",
  "Done when: new scenes for the pilot district are catalogued automatically.",
  10, ["I2", "D3"])
t("I8", B, c, "Open Satellite", "Copernicus Sentinel-1/2 and Landsat connector",
  "Open imagery for flood-extent and land-cover validation (no agreement needed).",
  "Done when: scenes over the pilot district are searchable by date and cloud cover.",
  8, ["I2"])
c = "5. Static & Reference Layers"
t("I9", B, c, "Terrain", "Ingest and tile SoI DEM",
  "Clean, reproject and tile the DEM; fill voids from open DEMs where needed.",
  "Done when: seamless DEM for the pilot district is served as tiles with void report.",
  6, ["I3", "D4"])
t("I10", B, c, "Population & Boundaries", "Census / LGD population and boundaries",
  "Load ward/village boundaries and population, households and age bands.",
  "Done when: every zone in the pilot district has population and household counts.",
  6, ["I3", "D5"])
t("I11", B, c, "Infrastructure", "Roads, bridges and critical facilities layer",
  "Build roads, bridges, hospitals, schools and govt buildings from OSM + Google Places/Roads.",
  "Done when: routable road network and facility layer cover the pilot district.",
  8, ["I3", "D6"])
t("I12", B, c, "History", "Historical disaster records",
  "Collect past flood events: dates, gauge peaks, mapped extents, damage and relief records.",
  "Done when: at least 5 past flood events are catalogued with extents for backtesting.",
  10, ["I3"], comment="Used to calibrate and backtest every model.")
t("I13", B, c, "Resources", "IDRN resource inventory import",
  "Import India Disaster Resource Network equipment and team inventory for the pilot state.",
  "Done when: IDRN inventory is queryable by location and resource type.",
  5, ["I3"], needs="Needs: IDRN data access via NDMA.")
c = "6. Sensor & Crowdsourced Feeds"
t("I14", B, c, "Hardware Sensors", "River / rain gauge sensor telemetry (pilot)",
  "Ingest telemetry from pilot river-level and rainfall gauges (MQTT/HTTP).",
  "Done when: pilot sensors stream into the store with device health status.",
  10, ["I2"], needs="Needs: sensor vendor / state agency feed access.")
t("I17", B, c, "Ground Reporting", "WhatsApp ground-reporting bot",
  "Structured bot for field staff and citizens to send photo + location + short description.",
  "Done when: reports arrive geotagged with consent recorded and reach the ingestion layer.",
  12, ["I2", "G3"], needs="Needs: WhatsApp Business API account.")
t("I20", B, c, "Ground Reporting", "Report verification and trust score",
  "De-duplicate, geo-check and score reports (image checks, reporter history, agreement with sensors) before they are trusted.",
  "Done when: every report carries a trust score and low-trust reports are held for review.",
  10, ["I17"], comment="Mitigation for crowdsourced misinformation risk.")
c = "7. Data Quality"
t("I21", B, c, "Normalisation", "Units, CRS and time-zone normalisation",
  "Convert all feeds to common units, one CRS and UTC/IST timestamps.",
  "Done when: no unit/CRS mismatches across Phase 1 sources in the QA report.",
  5, ["I5", "I6"])
t("I22", B, c, "Quality Rules", "Staleness, gap and outlier rules with alerts",
  "Rules that flag stale feeds, missing values and impossible readings; alert the data owner.",
  "Done when: rules are documented, applied and alerts fire in test.",
  6, ["I21"])

# ---------------------------------------------------------------- C
C = "C – AI/ML Prediction Engine (Layer 2)"
c = "8. EDA & Feature Prep"
t("M1", C, c, "EDA", "EDA on historical floods, rainfall and gauges",
  "Study rainfall–gauge lag, upstream/downstream relationships and past extents in the pilot basin.",
  "Done when: findings and modelling decisions are documented in a notebook.",
  8, ["I12", "I6", "I5"])
t("M2", C, c, "Features", "Feature store (rainfall accumulations, upstream levels, soil, land cover)",
  "Reusable, versioned features for all flood models.",
  "Done when: features are reproducible from raw data with one command.",
  8, ["M1", "I22"])
c = "9. Flood Forecasting"
t("M3", C, c, "Hydraulic Model", "Hydraulic/hydrological model set-up on DEM",
  "Set up a physics model (e.g. HEC-RAS 2D / LISFLOOD-FP) for the pilot basin and calibrate on past events.",
  "Done when: model reproduces at least 3 historical extents within the agreed accuracy bar.",
  15, ["I9", "M1"])
t("M4", C, c, "ML Forecaster", "Gauge-level forecaster (6/12/24/48 h)",
  "ML model (e.g. LSTM / gradient boosting) forecasting gauge levels from rainfall and upstream levels.",
  "Done when: beats CWC persistence baseline on held-out seasons.",
  15, ["M2"])
t("M5", C, c, "Flood Extent", "Flood-extent prediction (hydraulic + ML surrogate)",
  "Fast surrogate of the hydraulic model driven by forecast levels to produce extents in minutes.",
  "Done when: surrogate extents match the full model within the agreed tolerance and run in < 5 min.",
  15, ["M3", "M4"])
t("M6", C, c, "Fusion", "Fuse IMD, CWC and ISRO signals into one forecast",
  "Combine agency forecasts, sensor data and satellite-observed extents into one localised forecast.",
  "Done when: one fused forecast is produced per cycle with source contributions logged.",
  10, ["M5", "I7", "I14"])
c = "10. Risk Scoring"
t("M7", C, c, "Zone Risk", "Per-zone risk score",
  "Continuously updated risk score per zone (ward / hex grid) combining hazard and exposure.",
  "Done when: every zone has a score refreshed each forecast cycle.",
  8, ["M6", "I10"])
t("M8", C, c, "Time-Boxed Forecast", "Hour-by-hour hazard extents",
  "Time-boxed forecasts (e.g. extent X by hour 18, Y by hour 24) instead of static labels.",
  "Done when: hourly extents for the forecast window are published per cycle.",
  6, ["M7"])
t("M9", C, c, "Uncertainty", "Confidence bands for extents and scores",
  "Show likely / worst-case extents and a confidence level so operators know how far to trust each forecast.",
  "Done when: every forecast carries a confidence band and calibration is checked.",
  6, ["M8"])
t("M10", C, c, "Thresholds", "Critical-threshold and alert rules",
  "Rules that raise alerts and trigger pre-positioning when a zone crosses a critical score.",
  "Done when: thresholds are agreed with NDMA and alerts fire correctly in replay.",
  4, ["M9", "P5"])
c = "11. Model Validation"
t("M11", C, c, "Backtesting", "Backtest against historical events",
  "Replay past floods end to end and compare predicted vs observed extents, timings and gauges.",
  "Done when: backtest report meets the agreed acceptance bar for all catalogued events.",
  10, ["M8", "I12"], comment="Mitigation for model-trust risk.")
t("M12", C, c, "Reporting", "Model card and accuracy report",
  "Plain-language model card: data, method, accuracy by zone, known weaknesses.",
  "Done when: reviewed and accepted by NDMA technical committee.",
  4, ["M11"])

# ---------------------------------------------------------------- D
D = "D – Digital Twin Simulation (Layer 3)"
c = "12. 3D Base Twin"
t("T1", D, c, "Engine", "3D engine selection and spike",
  "Compare CesiumJS, deck.gl and similar for terrain, 3D tiles, time animation and offline support.",
  "Done when: engine chosen with a working spike on pilot terrain.",
  5, ["P3"])
t("T2", D, c, "Terrain", "Terrain mesh from DEM (3D tiles)",
  "Convert the DEM into streamable 3D terrain tiles.",
  "Done when: pilot district terrain loads smoothly in the browser.",
  8, ["T1", "I9"])
t("T3", D, c, "Infrastructure", "Buildings and infrastructure 3D layer",
  "Extrude buildings and add roads, bridges and critical facilities to the twin.",
  "Done when: critical facilities and road network are visible and selectable.",
  10, ["T2", "I11"])
t("T4", D, c, "Imagery", "Satellite imagery drape",
  "Drape latest optical imagery (open + commercial) over terrain.",
  "Done when: imagery layer can be switched by date.",
  5, ["T2", "I8"])
c = "13. Hazard Simulation"
t("T5", D, c, "Animation", "Animated flood spread over forecast window",
  "Render forecast extents and depths as an animated layer on the twin.",
  "Done when: the twin animates predicted spread hour by hour for any forecast cycle.",
  10, ["T3", "M8"])
t("T6", D, c, "Timeline", "Timeline scrubber (rewind / fast-forward)",
  "Rotate, zoom and scrub through the forecast timeline.",
  "Done when: operators can jump to any forecast hour in < 2 s.",
  6, ["T5"])
t("T7", D, c, "Impact View", "Cut-off roads and threatened neighbourhoods",
  "Highlight roads predicted to flood and neighbourhoods at risk, readable by non-technical users.",
  "Done when: cut-off roads and threatened areas are highlighted per forecast hour.",
  6, ["T5", "I11"])
t("T8", D, c, "Performance", "Performance on SEOC / district hardware",
  "Optimise tiles and rendering for typical control-room machines and bandwidth.",
  "Done when: first load < 5 s and smooth playback on agreed reference hardware.",
  5, ["T6"])
c = "14. Shadow Mode"
t("T9", D, c, "Replay", "Replay historical events in the twin",
  "Run the twin against past floods to validate behaviour before live use.",
  "Done when: any catalogued historical event can be replayed end to end.",
  8, ["T6", "M11"])
t("T10", D, c, "Comparison", "Predicted vs observed side-by-side",
  "Overlay observed extents (satellite) on predicted extents with an error summary.",
  "Done when: side-by-side view and error summary are available for every replayed event.",
  6, ["T9"])

# ---------------------------------------------------------------- E
E = "E – Planning & Decision Support (Layer 4)"
c = "15. Impact & Affected-Area Analysis"
t("E1", E, c, "Affected Population", "Affected population and households per zone",
  "Overlay Census/LGD population on hazard extents to estimate people and households affected.",
  "Done when: per-zone affected counts are produced for every forecast hour.",
  8, ["M7", "I10"])
t("E2", E, c, "Infrastructure", "Critical infrastructure exposure",
  "List hospitals, schools, bridges, substations etc. inside predicted extents.",
  "Done when: exposed assets are listed per zone and hour.",
  6, ["E1", "I11"])
t("E3", E, c, "Projections", "Time-boxed impact projections",
  "Impact numbers by hour (e.g. 'X households by hour 18, Y by hour 24').",
  "Done when: projections refresh automatically with each forecast cycle.",
  5, ["E1", "M8"])
t("E4", E, c, "Vulnerability", "Vulnerable-pocket flagging",
  "Flag high-risk zones with vulnerable groups (e.g. elderly) lacking nearby heavy-equipment teams via IDRN.",
  "Done when: vulnerable pockets are flagged with the reason shown.",
  8, ["E2", "I13"])
c = "16. Shelter & Safe-House Planning"
t("E5", E, c, "Shelter Candidates", "Candidate shelters outside the hazard zone",
  "Identify schools, community halls and govt buildings (Places/POI + local lists) outside predicted extents.",
  "Done when: candidate list updates with each forecast and excludes sites inside the extent.",
  8, ["E1", "I11"])
t("E6", E, c, "Shelter Candidates", "Shelter suitability scoring",
  "Score candidates on capacity, access route safety, distance and facilities.",
  "Done when: ranked shortlist is produced per affected zone.",
  6, ["E5"])
t("E7", E, c, "Capacity", "Live shelter capacity tracking",
  "Track occupancy so operators see which shelters are filling up.",
  "Done when: occupancy can be updated by field staff and shows on the dashboard.",
  8, ["E6"])
t("E8", E, c, "Capacity", "Shelter demand vs capacity gap",
  "Compare projected displaced population with available capacity by zone and hour.",
  "Done when: gaps are flagged with suggested extra sites.",
  5, ["E7", "E3"])
c = "17. Evacuation Route Planning"
t("E9", E, c, "Road Graph", "Routable road graph with predicted closures",
  "Mark road segments predicted to flood or close for each forecast hour.",
  "Done when: graph reflects predicted closures per forecast hour.",
  10, ["I11", "M8"])
t("E10", E, c, "Routing", "Safe evacuation route computation",
  "Compute evacuation routes from affected zones to shelters avoiding predicted closures.",
  "Done when: every affected zone has a safe route to its assigned shelter or is flagged as cut off.",
  8, ["E9", "E6"])
t("E11", E, c, "Routing", "Relief-vehicle routing with live traffic",
  "Routes for relief vehicles using Google Roads/Traffic plus predicted hazard.",
  "Done when: relief routes include live traffic and avoid predicted closures.",
  6, ["E10", "D6"])
t("E12", E, c, "Routing", "Automatic recompute on forecast update",
  "Recompute routes whenever the forecast changes so plans are always current.",
  "Done when: routes refresh within 5 min of a new forecast.",
  5, ["E10"])
t("E13", E, c, "De-confliction", "Citizen vs NDRF route conflict detection",
  "Flag where citizen evacuation routes and NDRF relief routes collide.",
  "Done when: conflicts are flagged with suggested alternatives.",
  6, ["E11"])
t("E14", E, c, "Export", "Export route and shelter plans for official channels",
  "Operator-issued export for SACHET / Cell Broadcast. Nothing is sent automatically.",
  "Done when: plans export in the agreed format after operator approval.",
  4, ["E12", "E8"], needs="Needs: SACHET / C-DOT message format confirmation.")

# ---------------------------------------------------------------- F
F = "F – Force Coordination & Control (Layer 5)"
c = "18. Force Tracking"
t("F1", F, c, "Unit Registry", "NDRF / SDRF unit and capability registry",
  "Register units, capabilities, equipment and home bases.",
  "Done when: all units for the pilot state are registered with capabilities.",
  5, ["I13"], needs="Needs: NDRF / SDRF unit data.")
t("F2", F, c, "Live Tracking", "GPS / radio tracking integration",
  "Plot responder positions from GPS devices or existing radio tracking.",
  "Done when: live positions update on the map at the agreed interval.",
  10, ["F1"], needs="Needs: NDRF tracking feed access.")
t("F3", F, c, "Live Tracking", "Forces on the 3D twin",
  "Show unit positions on the hazard simulation.",
  "Done when: units appear on the twin with status and last-seen time.",
  5, ["F2", "T6"])
c = "19. Predictive Pre-positioning"
t("F4", F, c, "Staging", "Staging locations outside predicted zone",
  "Choose staging points that remain outside the predicted extent and are reachable.",
  "Done when: staging candidates stay valid across the whole forecast window.",
  8, ["M10", "E5"])
t("F5", F, c, "Deployment Plans", "Auto-generated deployment plan",
  "When a zone crosses the critical threshold, recommend which units move, where and by when.",
  "Done when: plans are generated within 5 min of a threshold crossing.",
  10, ["F4", "F1"])
t("F6", F, c, "Deployment Plans", "ETA estimation",
  "Estimated arrival times using hazard-aware routing.",
  "Done when: every plan shows ETAs that account for predicted closures.",
  5, ["F5", "E11"])
t("F7", F, c, "Human in the Loop", "Commander approval workflow",
  "Commanders accept, edit or reject recommendations; all decisions are logged.",
  "Done when: no recommendation is actioned without approval and every decision is audited.",
  5, ["F5"])
c = "20. Unified Command View"
t("F8", F, c, "SEOC Screen", "Unified SEOC command screen",
  "One screen with hazard zones, force deployments and ETAs, replacing radio check-ins.",
  "Done when: SEOC staff can run a full scenario from this one screen.",
  8, ["F3", "F6", "F7", "E8"])

# ---------------------------------------------------------------- G
G = "G – Command Dashboard & Feedback Loop"
c = "21. Dashboard UX"
t("U1", G, c, "UX Research", "Wireframes from user workflows",
  "Low-fidelity wireframes of all panels based on the workflow study.",
  "Done when: wireframes are reviewed with NDMA and SEOC users.",
  10, ["P4"])
t("U2", G, c, "Design", "Design system and high-fidelity mockups",
  "Visual design, colour scales for risk, accessibility and control-room (large screen) layouts.",
  "Done when: mockups approved by NDMA.",
  8, ["U1"])
c = "22. Dashboard Build"
t("U3", G, c, "Shell", "Dashboard shell (login, layout, layer toggles)",
  "Base app: SSO login, layout, map/twin container and layer controls.",
  "Done when: shell is deployed in test with SSO.",
  10, ["U2", "T1", "H3"])
t("U4", G, c, "Shell", "Role-based views (NDMA, SEOC, DEOC, NDRF)",
  "Each role sees the panels and areas relevant to it.",
  "Done when: each role sees only its permitted views and data.",
  6, ["U3"])
t("U5", G, c, "Panels", "Risk and alerts panel",
  "Per-zone risk scores, threshold alerts and confidence.",
  "Done when: alerts appear within one forecast cycle of a threshold crossing.",
  6, ["U3", "M10"])
t("U6", G, c, "Panels", "Impact panel",
  "Affected population, households, infrastructure and vulnerable pockets.",
  "Done when: panel matches the impact analysis outputs.",
  5, ["U3", "E4"])
t("U7", G, c, "Panels", "Shelter panel",
  "Shelter recommendations, live capacity and gaps.",
  "Done when: operators can review, edit and lock a shelter plan.",
  5, ["U3", "E8"])
t("U8", G, c, "Panels", "Route planning workspace",
  "Evacuation and relief routes with conflicts and export.",
  "Done when: operators can generate, compare and export route plans.",
  6, ["U3", "E13", "E14"])
t("U9", G, c, "Accessibility", "English / Hindi interface",
  "Bilingual labels and reports for district users.",
  "Done when: all labels and exports are available in English and Hindi.",
  4, ["U4"])
t("U10", G, c, "Accessibility", "Low-bandwidth mode for district EOCs",
  "Lighter 2D fallback and cached tiles for poor connectivity.",
  "Done when: dashboard is usable at the agreed low-bandwidth profile.",
  6, ["U4", "T8"])
c = "23. Continuous Feedback"
t("FB1", G, c, "Feedback Loop", "Ground reports update the forecast",
  "Verified ground reports adjust the observed state and trigger a re-forecast.",
  "Done when: a verified report changes the next forecast cycle and is traceable.",
  8, ["I20", "M8"])
t("FB2", G, c, "Feedback Loop", "Sensor and force updates in near real time",
  "Live sensor and force-position updates refresh forecast, routes and plans.",
  "Done when: dashboard refreshes within the agreed latency after new data.",
  8, ["FB1", "F2", "E12"])
t("FB3", G, c, "Feedback Loop", "Operator verification queue for reports",
  "Queue where operators confirm or reject low-trust reports.",
  "Done when: every low-trust report is reviewed or expires with a reason.",
  5, ["I20", "U3"])

# ---------------------------------------------------------------- H
H = "H – Platform, Security & MLOps"
c = "24. Infrastructure"
t("H0", H, c, "Hosting", "Hosting decision (MeghRaj / NIC / GovCloud)",
  "Choose a government-approved cloud with data residency in India.",
  "Done when: hosting is approved in writing and accounts are provisioned.",
  5, ["P1"], needs="Needs: NDMA IT approval.")
t("H1", H, c, "Environments", "Dev / test / prod environments as code",
  "Infrastructure-as-code for all environments.",
  "Done when: environments can be rebuilt from code.",
  8, ["H0"])
t("H2", H, c, "Delivery", "CI/CD pipeline",
  "Automated tests, builds and deployments.",
  "Done when: every merge deploys to test automatically.",
  5, ["H1"])
t("H10", H, c, "Resilience", "Backup and disaster recovery",
  "Backups and a recovery plan — the system must stay up during disasters.",
  "Done when: restore drill succeeds within the agreed RTO/RPO.",
  5, ["H1"])
c = "25. Security"
t("H3", H, c, "Access", "SSO, role-based access and audit logs",
  "Single sign-on, role-based permissions and audit trail of every action.",
  "Done when: all actions are attributable to a user in the audit log.",
  6, ["H1"])
t("H5", H, c, "Data Protection", "Encryption and secrets management",
  "Encrypt data at rest/in transit; keep keys in a secrets manager.",
  "Done when: no plaintext secrets in code or config; encryption verified.",
  4, ["H1"])
t("H4", H, c, "Audit", "Security audit (CERT-In empanelled auditor)",
  "Third-party security audit before any live use.",
  "Done when: audit passed with all high/critical findings closed.",
  10, ["V1"], needs="Needs: auditor engagement.")
c = "26. MLOps"
t("H6", H, c, "Registry", "Model registry and versioning",
  "Track every model version, data snapshot and config.",
  "Done when: every forecast is traceable to a model version.",
  5, ["H1", "M4"])
t("H7", H, c, "Retraining", "Scheduled retraining and re-calibration",
  "Retrain on new seasons/events and promote only if it beats the current model.",
  "Done when: retraining runs on schedule and promotion is gated by backtest scores.",
  8, ["H6", "M11"])
t("H8", H, c, "Monitoring", "Monitoring (feeds, latency, failures, drift)",
  "Dashboards and alerts for feed staleness, forecast latency, failures and model drift.",
  "Done when: on-call is alerted for every agreed failure mode.",
  8, ["H6", "I22"])
t("H9", H, c, "Load", "Load test for disaster-surge usage",
  "Test with surge traffic and data volumes expected during a live event.",
  "Done when: system meets latency targets at 3× expected peak.",
  5, ["F8"])

# ---------------------------------------------------------------- I
V = "I – Pilot, Validation & Rollout"
c = "27. Testing"
t("V1", V, c, "Integration", "End-to-end integration testing",
  "Run the full chain: feeds → forecast → twin → impact → shelters/routes → force plans.",
  "Done when: all end-to-end test scenarios pass in the test environment.",
  10, ["F8", "U5", "U6", "U7", "U8", "FB2", "FB3", "H8"])
t("V2", V, c, "Tabletop", "Tabletop exercise with SEOC (simulated flood)",
  "Run a simulated flood with SEOC and DEOC staff using the dashboard.",
  "Done when: exercise completed and issues logged and triaged.",
  5, ["V1"])
t("V3", V, c, "UAT", "User acceptance testing with NDMA / SEOC",
  "Formal UAT against the agreed success metrics.",
  "Done when: UAT sign-off received.",
  10, ["V2", "H4"])
c = "28. Training & Handover"
t("V6", V, c, "Training", "Operator training and SOPs",
  "Train NDMA, SEOC and DEOC operators; write SOPs for each workflow.",
  "Done when: all pilot operators trained and SOPs approved.",
  8, ["V3"])
t("V7", V, c, "Documentation", "Technical documentation and runbooks",
  "Architecture, data and model docs; on-call runbooks.",
  "Done when: runbooks tested by someone outside the build team.",
  6, ["V3"])
c = "29. Shadow-Mode Pilot"
t("V4", V, c, "Shadow Mode", "Shadow run through monsoon 2027",
  "Run the twin alongside existing processes in the pilot district; decisions are not driven by it yet.",
  "Done when: monsoon season run complete with forecast-vs-actual log for every event.",
  58, ["V6", "V7", "T10", "M12"], not_before=date(2027, 6, 14),
  comment="Covers peak monsoon (Jul–Sep 2027). Proposal requires shadow mode before the twin drives live decisions.")
t("V5", V, c, "Go / No-Go", "Go / no-go review for live decision support",
  "Review shadow-mode results with NDMA against the acceptance bar.",
  "Done when: written go/no-go decision for live use.",
  5, ["V4"])
c = "30. Phase 2 & 3 Expansion"
t("X1", V, c, "Phase 2 – Hazards", "Landslide slope-stability model",
  "Rainfall-threshold and slope-stability model using soil moisture sensors.",
  "Done when: backtested on past landslides in a hill pilot district.",
  25, ["V5"])
t("X2", V, c, "Phase 2 – Hazards", "Cyclone track and surge impact model",
  "Use IMD cyclone tracks for wind and storm-surge impact in a coastal district.",
  "Done when: backtested on past cyclones.",
  25, ["V5"])
t("X3", V, c, "Phase 2 – Inputs", "Soil moisture and structural health sensors",
  "Ingest soil moisture (slopes) and vibration/tilt (bridges, dams) telemetry.",
  "Done when: sensors stream with health checks and feed the landslide model.",
  10, ["V5"])
t("X4", V, c, "Phase 2 – Inputs", "CCTV / traffic camera water-level estimation",
  "Computer vision on camera feeds to estimate flood depth and road conditions.",
  "Done when: CV estimates validated against gauges at test sites.",
  15, ["V5"])
t("X5", V, c, "Phase 2 – Inputs", "Social media monitoring with NLP filtering",
  "Geotagged posts (X, Facebook, Instagram) filtered by NLP before being trusted.",
  "Done when: event-detection precision meets the agreed bar on past events.",
  15, ["V5"], needs="Needs: platform API access and privacy approval.")
t("X6", V, c, "Phase 2 – Inputs", "SACHET app ground-reporting extension",
  "First-party ground reporting inside SACHET, feeding the same ingestion layer.",
  "Done when: SACHET reports arrive in the ingestion layer with consent.",
  15, ["V5"], needs="Needs: NDMA / C-DOT agreement for SACHET changes.")
t("X7", V, c, "Phase 3 – Scale-up", "Multi-district / state rollout plan",
  "Onboard additional districts using lessons from the pilot; DEM and sensor gap analysis per district.",
  "Done when: rollout plan with district order, data gaps and costs approved.",
  15, ["X1", "X2"])

# ------------------------------------------------------------------- schedule

def is_work(d):
    return d.weekday() < 5 and d not in HOL


def next_work(d):
    while not is_work(d):
        d += timedelta(days=1)
    return d


def add_work(start, n):
    d, left = start, n - 1
    while left > 0:
        d += timedelta(days=1)
        if is_work(d):
            left -= 1
    return d


by_id = {r["id"]: r for r in R}
done = set()


def sched(r):
    if r["id"] in done:
        return
    for dep in r["deps"]:
        sched(by_id[dep])
    if r["start"] is None:
        s = PROJECT_START
        for dep in r["deps"]:
            s = max(s, by_id[dep]["end"] + timedelta(days=1))
        if r["not_before"]:
            s = max(s, r["not_before"])
        r["start"] = next_work(s)
    else:
        r["start"] = next_work(r["start"])
    r["end"] = add_work(r["start"], r["dur"])
    done.add(r["id"])


for r in R:
    sched(r)


def short(rid):
    return f"{by_id[rid]['sub']} ({rid})"


for r in R:
    parts = []
    if r["deps"]:
        parts.append("After: " + "; ".join(short(d) for d in r["deps"]) + ".")
    if r["needs"]:
        parts.append(r["needs"])
    r["depend"] = " ".join(parts)

# ------------------------------------------------------------------- styles
NAVY = "193F7F"
MOD_FILL = PatternFill("solid", fgColor="D8E8F9")
GREY_FILL = PatternFill("solid", fgColor="F4F4F4")
NAVY_FILL = PatternFill("solid", fgColor=NAVY)
WHITE = PatternFill("solid", fgColor="FFFFFF")
FONT = "Calibri"
f_title = Font(name=FONT, size=16, bold=True, color="FFFFFF")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_bold = Font(name=FONT, size=10, bold=True)
f_body = Font(name=FONT, size=10)
f_note = Font(name=FONT, size=10, italic=True, color="595959")
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
DATE_FMT = "dd mmm yy"

STATUS_STYLES = {
    "Done": ("B7E0CC", "0C6633"),
    "In Progress": ("FFEAB2", "996600"),
    "Not Started": ("E0E0E0", "595959"),
    "Blocked": ("F4CCCC", "B21919"),
}

wb = Workbook()
ws = wb.active
ws.title = "Tracker"
ws.sheet_view.showGridLines = False
HEAD = ["Module", "Category", "Task", "Sub Task", "Description",
        "Acceptance criteria / Definition of done", "Start", "End", "Working days",
        "Status", "Dependency / Blocker", "Comments", "Artifacts"]
WIDTHS = [150, 135, 135, 230, 330, 330, 80, 80, 80, 105, 230, 230, 230]
for i, px in enumerate(WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = round(px / 7, 1)

ws.merge_cells("A1:M1")
ws["A1"] = "NDMA — AI-Integrated Digital Twin Command Dashboard · Project Tracker"
ws["A1"].font = f_title
ws["A1"].fill = NAVY_FILL
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 30
ws.merge_cells("A2:M2")
ws["A2"] = ("Source: NDMA_Digital_Twin_Proposal.pdf · Phase 1 = one pilot district, riverine flood · "
            "Phase 2 = landslide, cyclone, more sensor/crowd inputs · Phase 3 = multi-district scale-up · "
            "Edit Start, End and Status; Working days recalculates (weekends and the Holidays tab excluded). "
            "Dates are planned from dependencies starting 12 Oct 26.")
ws["A2"].font = f_note
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
ws.row_dimensions[2].height = 30
for i, h in enumerate(HEAD, 1):
    cell = ws.cell(3, i, h)
    cell.font = f_head
    cell.fill = NAVY_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    cell.border = BORDER
ws.row_dimensions[3].height = 28
ws.freeze_panes = "A4"

FIRST = 4
row = FIRST
for r in R:
    r["row"] = row
    vals = [None, None, None, r["sub"], r["desc"], r["dod"], r["start"], r["end"],
            f"=NETWORKDAYS(G{row},H{row},Holidays!$A$2:$A$40)", r["status"], r["depend"],
            r["comment"], r["artifact"]]
    for i, v in enumerate(vals, 1):
        cell = ws.cell(row, i, v)
        cell.font = f_body
        cell.fill = WHITE
        cell.border = BORDER
        cell.alignment = WRAP_TOP
    for col in (7, 8):
        ws.cell(row, col).number_format = DATE_FMT
        ws.cell(row, col).alignment = CENTER_TOP
    ws.cell(row, 9).alignment = CENTER_TOP
    st = ws.cell(row, 10)
    st.font = f_bold
    st.alignment = CENTER_TOP
    row += 1
LAST = row - 1


def merge_groups(col, key, fill, bold=True):
    start = FIRST
    for i in range(1, len(R) + 1):
        end_group = i == len(R) or key(R[i]) != key(R[start - FIRST])
        if end_group:
            first, last = start, FIRST + i - 1
            cell = ws.cell(first, col, R[start - FIRST][{1: "mod", 2: "cat", 3: "task"}[col]])
            for rr in range(first, last + 1):
                ws.cell(rr, col).fill = fill
                ws.cell(rr, col).border = BORDER
            cell.font = f_bold if bold else f_body
            cell.alignment = WRAP_TOP
            if last > first:
                ws.merge_cells(start_row=first, start_column=col, end_row=last, end_column=col)
            start = FIRST + i


merge_groups(1, lambda r: r["mod"], MOD_FILL)
merge_groups(2, lambda r: (r["mod"], r["cat"]), GREY_FILL)
merge_groups(3, lambda r: (r["mod"], r["cat"], r["task"]), GREY_FILL)

dv = DataValidation(type="list", formula1='"Done,In Progress,Not Started,Blocked"', allow_blank=False,
                    showErrorMessage=True, errorTitle="Status", error="Pick a status from the list.")
ws.add_data_validation(dv)
dv.add(f"J{FIRST}:J{LAST + 60}")
for name, (bg, fg) in STATUS_STYLES.items():
    ws.conditional_formatting.add(
        f"J{FIRST}:J{LAST + 60}",
        FormulaRule(formula=[f'$J{FIRST}="{name}"'],
                    fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                    font=Font(bold=True, color=fg)))
# Overdue: end date passed and not done.
ws.conditional_formatting.add(
    f"H{FIRST}:H{LAST + 60}",
    FormulaRule(formula=[f'AND($H{FIRST}<>"",$H{FIRST}<TODAY(),$J{FIRST}<>"Done")'],
                font=Font(bold=True, color="B21919")))

# ------------------------------------------------------------------- Summary
sm = wb.create_sheet("Summary")
sm.sheet_view.showGridLines = False
sm.merge_cells("A1:J1")
sm["A1"] = "Summary by Module"
sm["A1"].font = f_title
sm["A1"].fill = NAVY_FILL
sm["A1"].alignment = Alignment(vertical="center", indent=1)
sm.row_dimensions[1].height = 30
sm["A2"] = "All figures are formulas over the Tracker tab; they update as Status and dates change."
sm["A2"].font = f_note
sh = ["Module", "Sub tasks", "Done", "In Progress", "Not Started", "Blocked", "% Done",
      "Start", "End", "Working days"]
for i, h in enumerate(sh, 1):
    cell = sm.cell(3, i, h)
    cell.font = f_head
    cell.fill = NAVY_FILL
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([42, 10, 10, 12, 12, 10, 10, 12, 12, 13], 1):
    sm.column_dimensions[get_column_letter(i)].width = w

mods = []
for r in R:
    if not mods or mods[-1][0] != r["mod"]:
        mods.append([r["mod"], r["row"], r["row"]])
    else:
        mods[-1][2] = r["row"]
srow = 4
for name, a, b in mods:
    rng = lambda col: f"Tracker!${col}${a}:${col}${b}"
    vals = [name, f"=COUNTA({rng('D')})",
            f'=COUNTIF({rng("J")},"Done")', f'=COUNTIF({rng("J")},"In Progress")',
            f'=COUNTIF({rng("J")},"Not Started")', f'=COUNTIF({rng("J")},"Blocked")',
            f"=IF(B{srow}=0,0,C{srow}/B{srow})", f"=MIN({rng('G')})", f"=MAX({rng('H')})",
            f"=NETWORKDAYS(H{srow},I{srow},Holidays!$A$2:$A$40)"]
    for i, v in enumerate(vals, 1):
        cell = sm.cell(srow, i, v)
        cell.font = f_bold if i == 1 else f_body
        cell.border = BORDER
        cell.fill = MOD_FILL if i == 1 else WHITE
    sm.cell(srow, 7).number_format = "0%"
    sm.cell(srow, 8).number_format = DATE_FMT
    sm.cell(srow, 9).number_format = DATE_FMT
    srow += 1
tot = srow
sm.cell(tot, 1, "Total").font = f_bold
for col in "BCDEF":
    sm[f"{col}{tot}"] = f"=SUM({col}4:{col}{tot - 1})"
sm[f"G{tot}"] = f"=IF(B{tot}=0,0,C{tot}/B{tot})"
sm[f"H{tot}"] = f"=MIN(H4:H{tot - 1})"
sm[f"I{tot}"] = f"=MAX(I4:I{tot - 1})"
sm[f"J{tot}"] = f"=NETWORKDAYS(H{tot},I{tot},Holidays!$A$2:$A$40)"
for i in range(1, 11):
    cell = sm.cell(tot, i)
    cell.font = f_bold
    cell.fill = GREY_FILL
    cell.border = BORDER
sm[f"G{tot}"].number_format = "0%"
sm[f"H{tot}"].number_format = DATE_FMT
sm[f"I{tot}"].number_format = DATE_FMT

# Key milestones (dates linked to the Tracker so they move with it).
ms_row = tot + 2
sm.cell(ms_row, 1, "Key milestones").font = Font(name=FONT, size=12, bold=True, color=NAVY)
for i, h in enumerate(["Milestone", "Tracker row", "Target date"], 1):
    cell = sm.cell(ms_row + 1, i, h)
    cell.font = f_head
    cell.fill = NAVY_FILL
    cell.border = BORDER
milestones = [
    ("NDMA sign-off on proposal", "P1"), ("Pilot district selected", "P2"),
    ("Core agency feeds live (IMD, CWC)", "I6"), ("Flood forecast backtest passed", "M11"),
    ("Digital twin animates forecast", "T6"), ("Unified SEOC command screen", "F8"),
    ("End-to-end integration test", "V1"), ("UAT sign-off", "V3"),
    ("Shadow mode complete (monsoon 2027)", "V4"), ("Go / no-go for live use", "V5"),
    ("Phase 3 rollout plan approved", "X7"),
]
mr = ms_row + 2
for label, rid in milestones:
    rr = by_id[rid]["row"]
    sm.cell(mr, 1, label).font = f_body
    sm.cell(mr, 2, f"=ROW(Tracker!H{rr})").font = f_body
    c_ = sm.cell(mr, 3, f"=Tracker!H{rr}")
    c_.font = Font(name=FONT, size=10, color="008000")
    c_.number_format = DATE_FMT
    for i in range(1, 4):
        sm.cell(mr, i).border = BORDER
    mr += 1

# ------------------------------------------------------------------- Decisions & dependencies
dd = wb.create_sheet("Decisions & Dependencies")
dd.sheet_view.showGridLines = False
dd.merge_cells("A1:G1")
dd["A1"] = "External Decisions & Dependencies"
dd["A1"].font = f_title
dd["A1"].fill = NAVY_FILL
dd["A1"].alignment = Alignment(vertical="center", indent=1)
dd.row_dimensions[1].height = 30
dd["A2"] = "Everything the team needs from NDMA or another agency. 'Needed by' = start of the first task that waits on it (linked to the Tracker)."
dd["A2"].font = f_note
dh = ["#", "Decision / dependency", "Owner", "Blocks", "Needed by", "Status", "Notes"]
for i, h in enumerate(dh, 1):
    cell = dd.cell(3, i, h)
    cell.font = f_head
    cell.fill = NAVY_FILL
    cell.border = BORDER
for i, w in enumerate([5, 48, 22, 40, 12, 13, 44], 1):
    dd.column_dimensions[get_column_letter(i)].width = w
decisions = [
    ("Proposal walkthrough and sign-off", "NDMA", ["P1"], ""),
    ("Nomination of candidate pilot districts", "NDMA / State SDMA", ["P2"], "Prefer a well-mapped, flood-prone district (terrain-resolution risk)."),
    ("Hosting / cloud approval (data residency in India)", "NDMA IT", ["H0"], "Blocks all environments."),
    ("IMD data-sharing agreement", "IMD", ["D1"], "Formal agreements from day one (inter-agency risk)."),
    ("CWC data-sharing agreement", "CWC", ["D2"], ""),
    ("ISRO / NRSC imagery access", "NRSC", ["D3"], ""),
    ("Survey of India DEM licence", "Survey of India", ["D4"], ""),
    ("Google Maps Platform budget", "NDMA Finance", ["D6"], ""),
    ("Commercial imagery procurement", "NDMA Procurement", ["D7"], "Phase 1 pilot district only."),
    ("Privacy policy legal review (DPDP Act 2023)", "NDMA Legal", ["G1"], ""),
    ("IDRN inventory access", "NDMA / NIDM", ["I13"], ""),
    ("Pilot sensor feed access", "State water resources dept / vendor", ["I14"], ""),
    ("WhatsApp Business API account", "NDMA", ["I17"], ""),
    ("NDRF / SDRF unit data and tracking feed", "NDRF / SDRF", ["F1", "F2"], ""),
    ("SACHET / Cell Broadcast export format", "NDMA / C-DOT", ["E14"], "Operators issue messages; system never auto-sends."),
    ("Security auditor engagement", "NDMA IT", ["H4"], "CERT-In empanelled."),
]
drow = 4
for n, (label, owner, blocks, note) in enumerate(decisions, 1):
    first = min(blocks, key=lambda b: by_id[b]["start"])
    vals = [n, label, owner, "; ".join(short(b) for b in blocks),
            f"=Tracker!G{by_id[first]['row']}", "Open", note]
    for i, v in enumerate(vals, 1):
        cell = dd.cell(drow, i, v)
        cell.font = f_body
        cell.border = BORDER
        cell.alignment = WRAP_TOP
    dd.cell(drow, 5).number_format = DATE_FMT
    dd.cell(drow, 5).font = Font(name=FONT, size=10, color="008000")
    drow += 1
dv2 = DataValidation(type="list", formula1='"Open,In Progress,Closed"', allow_blank=False)
dd.add_data_validation(dv2)
dv2.add(f"F4:F{drow + 20}")
dd.freeze_panes = "A4"

# ------------------------------------------------------------------- Risks
rk = wb.create_sheet("Risks")
rk.sheet_view.showGridLines = False
rk.merge_cells("A1:G1")
rk["A1"] = "Risk Register"
rk["A1"].font = f_title
rk["A1"].fill = NAVY_FILL
rk["A1"].alignment = Alignment(vertical="center", indent=1)
rk.row_dimensions[1].height = 30
rk["A2"] = "Risks 1–5 from proposal section 7; 6–8 added for delivery. Impact / Likelihood are initial team estimates — review with NDMA."
rk["A2"].font = f_note
rh = ["#", "Risk", "Impact", "Likelihood", "Mitigation", "Linked tasks", "Status"]
for i, h in enumerate(rh, 1):
    cell = rk.cell(3, i, h)
    cell.font = f_head
    cell.fill = NAVY_FILL
    cell.border = BORDER
for i, w in enumerate([5, 44, 10, 11, 60, 34, 10], 1):
    rk.column_dimensions[get_column_letter(i)].width = w
risks = [
    ("Terrain data resolution — inconsistent elevation data distorts hazard-flow predictions.", "High", "Medium",
     "Launch Phase 1 in a single well-mapped pilot district before scaling.", "P2, D4, I9"),
    ("Crowdsourced data reliability — misinformation or false reports corrupt the model.", "High", "High",
     "Strict NLP/verification filtering and trust scores before ingestion; operator review queue.", "I20, FB3, X5"),
    ("Model trust & validation — wrong predictions misdirect forces or plans.", "High", "Medium",
     "Backtest against historical data and run in shadow mode before live decisions.", "M11, T9, T10, V4, V5"),
    ("Data privacy & consent — citizen location and social data create legal liability.", "High", "Medium",
     "Data-governance and privacy policy built alongside the technical build; consent on every report.", "G1, G2, G3"),
    ("Inter-agency coordination — IMD, CWC, ISRO use siloed, incompatible formats.", "High", "High",
     "Formal data-sharing agreements and a common schema from day one.", "D1, D2, D3, I1"),
    ("Agreement delays push the schedule (agency approvals are outside team control).", "High", "High",
     "Start agreements immediately after pilot selection; build on open data (Copernicus, OSM) meanwhile.", "D1–D7, I8"),
    ("Shadow-mode window depends on monsoon timing; a quiet season gives few events.", "Medium", "Medium",
     "Supplement with historical replays; extend shadow mode or add a second pilot basin if needed.", "V4, T9"),
    ("System unavailable during a live disaster (surge load, outages).", "High", "Low",
     "Load testing at 3× peak, DR drill and low-bandwidth mode.", "H9, H10, U10"),
]
rrow = 4
for n, (risk, imp, lik, mit, links) in enumerate(risks, 1):
    for i, v in enumerate([n, risk, imp, lik, mit, links, "Open"], 1):
        cell = rk.cell(rrow, i, v)
        cell.font = f_body
        cell.border = BORDER
        cell.alignment = WRAP_TOP
    rrow += 1
rk.freeze_panes = "A4"

# ------------------------------------------------------------------- Holidays
hs = wb.create_sheet("Holidays")
hs["A1"], hs["B1"] = "Date", "Holiday"
for col in "AB":
    hs[f"{col}1"].font = f_head
    hs[f"{col}1"].fill = NAVY_FILL
hs.column_dimensions["A"].width = 14
hs.column_dimensions["B"].width = 40
for i, (d, name) in enumerate(HOLIDAYS, 2):
    hs.cell(i, 1, d).number_format = DATE_FMT
    hs.cell(i, 2, name)
hs["D1"] = ("Edit this list to match the NDMA gazetted holiday calendar; Working days on every tab "
            "uses A2:A40. The planned dates on the Tracker were scheduled with this list.")
hs["D1"].font = f_note
hs["D1"].alignment = Alignment(wrap_text=True, vertical="top")
hs.column_dimensions["D"].width = 60

wb.save(OUT)
print(f"rows={len(R)} first={FIRST} last={LAST}")
print("end of plan:", max(r["end"] for r in R))
for rid in ["P1", "P2", "D1", "I6", "M11", "T6", "F8", "V1", "V3", "V4", "V5", "X7"]:
    r = by_id[rid]
    print(rid, r["start"], r["end"])
